# -*- coding: utf-8 -*-
"""Gera o template de projeto de pesquisa quantitativa (DOCX), padrao UFMA."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

VINHO = RGBColor(0x8D, 0x03, 0x33)
DOURADO = RGBColor(0xD4, 0xB2, 0x77)
CINZA = RGBColor(0x55, 0x55, 0x55)
PRETO = RGBColor(0x1A, 0x1A, 0x1A)

F_TIT = "Montserrat"
F_TXT = "Source Sans Pro"

# O logotipo oficial da UFMA, se disponivel localmente. Sem ele o cabecalho sai
# so com o bloco de texto institucional, e o documento continua valido.
RAIZ = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
LOGO = os.path.join(RAIZ, "slides", "assets", "logo-ufma.png")
if not os.path.exists(LOGO):
    LOGO = os.environ.get("LOGO_UFMA", "")

CONTENT_W = Cm(16.0)

doc = Document()

# ---------------------------------------------------------------- pagina e estilos
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0)
sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)

normal = doc.styles["Normal"]
normal.font.name = F_TXT
normal.font.size = Pt(11)
normal.font.color.rgb = PRETO
normal._element.rPr.rFonts.set(qn("w:eastAsia"), F_TXT)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
pf.line_spacing = 1.15
pf.space_before = Pt(0)
pf.space_after = Pt(8)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _shd(el, hexcolor):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    el.append(shd)


def _pborder(p, hexcolor="8D0333", sz=6, sides=("bottom",)):
    pPr = p._p.get_or_add_pPr()
    bd = pPr.find(qn("w:pBdr"))
    if bd is None:
        bd = OxmlElement("w:pBdr")
        pPr.append(bd)
    for s in sides:
        e = OxmlElement("w:" + s)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "4")
        e.set(qn("w:color"), hexcolor)
        bd.append(e)


def _keep_next(p):
    pPr = p._p.get_or_add_pPr()
    e = OxmlElement("w:keepNext")
    pPr.append(e)


def _partes(item):
    """Normaliza: uma tupla (texto, bold[, italic]) vira lista de uma parte."""
    if isinstance(item, str):
        return [(item, False)]
    if isinstance(item, (list, tuple)) and len(item) >= 2 and isinstance(item[1], bool):
        return [item]
    return list(item)


def run(p, text, size=11, bold=False, italic=False, color=PRETO, font=F_TXT, caps=False):
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.all_caps = caps
    return r


def para(text="", size=11, bold=False, italic=False, color=PRETO, font=F_TXT,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8, indent=None, left=None):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent is not None:
        p.paragraph_format.first_line_indent = indent
    if left is not None:
        p.paragraph_format.left_indent = left
    if text:
        run(p, text, size=size, bold=bold, italic=italic, color=color, font=font)
    return p


def rich(parts, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8, left=None):
    """parts: lista de (texto, bold) ou (texto, bold, italic)."""
    p = doc.add_paragraph()
    p.paragraph_format.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if left is not None:
        p.paragraph_format.left_indent = left
    for t in parts:
        txt, bold = t[0], t[1]
        it = t[2] if len(t) > 2 else False
        run(p, txt, size=size, bold=bold, italic=it)
    return p


def h1(num, text):
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    p.style = doc.styles["Heading 1"]
    run(p, (num + "  " if num else "") + text, size=14, bold=True, color=VINHO, font=F_TIT)
    _pborder(p, "D4B277", 8)
    _keep_next(p)
    return p


def h1_nobreak(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    p.style = doc.styles["Heading 1"]
    run(p, (num + "  " if num else "") + text, size=14, bold=True, color=VINHO, font=F_TIT)
    _pborder(p, "D4B277", 8)
    _keep_next(p)
    return p


def h2(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.style = doc.styles["Heading 2"]
    run(p, (num + "  " if num else "") + text, size=12, bold=True, color=VINHO, font=F_TIT)
    _keep_next(p)
    return p


def h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.style = doc.styles["Heading 3"]
    run(p, text, size=11, bold=True, color=PRETO, font=F_TIT)
    _keep_next(p)
    return p


def bullets(items, size=11):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for t in _partes(it):
            txt, bold = t[0], t[1]
            it2 = t[2] if len(t) > 2 else False
            run(p, txt, size=size, bold=bold, italic=it2)


def numbers(items, size=11):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for t in _partes(it):
            txt, bold = t[0], t[1]
            it2 = t[2] if len(t) > 2 else False
            run(p, txt, size=size, bold=bold, italic=it2)


def _one_cell_box(fill, border, pad_left=Cm(0.25)):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    t.columns[0].width = CONTENT_W
    c = t.cell(0, 0)
    c.width = CONTENT_W
    tcPr = c._tc.get_or_add_tcPr()
    if fill:
        _shd(tcPr, fill)
    borders = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + s)
        if border:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "6")
            e.set(qn("w:color"), border)
        else:
            e.set(qn("w:val"), "nil")
        borders.append(e)
    tcPr.append(borders)
    mar = OxmlElement("w:tcMar")
    for s, v in (("top", 80), ("left", 120), ("bottom", 80), ("right", 120)):
        e = OxmlElement("w:" + s)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)
    return t, c


def caixa_info(titulo, partes, fill="F6EEF2", border="C88BA0"):
    """Caixa de instrucao. `partes` e uma lista de str (texto corrido) e de
    tuplas (texto, bold[, italic]); uma str vazia abre novo paragrafo."""
    t, c = _one_cell_box(fill, border)
    c.paragraphs[0].text = ""

    def novo(primeiro=False):
        p = c.paragraphs[0] if primeiro else c.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p

    p = novo(primeiro=True)
    if titulo:
        run(p, titulo + " ", size=9.5, bold=True, color=VINHO, font=F_TIT)
    for parte in partes:
        if isinstance(parte, (list, tuple)):
            txt, bold = parte[0], parte[1]
            it2 = parte[2] if len(parte) > 2 else False
            run(p, txt, size=9.5, bold=bold, italic=it2, color=CINZA)
        elif parte == "":
            p = novo()
        else:
            run(p, parte, size=9.5, color=CINZA)
    espaco()
    return t


def caixa_exemplo(linhas, titulo="Exemplo preenchido."):
    caixa_info(titulo, linhas, fill="FBF6EC", border="D4B277")


def espaco(pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run(p, "", size=pts)
    return p


def campo(rotulo=None, linhas=3, dica=None):
    """Caixa em branco para o estudante escrever."""
    if rotulo:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run(p, rotulo, size=10.5, bold=True, color=VINHO, font=F_TIT)
        if dica:
            run(p, "   " + dica, size=9, italic=True, color=CINZA)
        _keep_next(p)
    t, c = _one_cell_box(None, "BFBFBF")
    c.paragraphs[0].text = ""
    first = True
    for _ in range(linhas):
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.5
        run(p, "", size=11)
    espaco()
    return t


def tabela(headers, rows, widths, fonte=None, legenda=None, header_fill="8D0333",
           size=9.5, header_size=9.5):
    if legenda:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run(p, legenda, size=10, bold=True, color=PRETO, font=F_TIT)
        _keep_next(p)
    total = sum(widths)
    scale = CONTENT_W.cm / total
    widths = [Cm(w * scale) for w in widths]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = w
    hdr = t.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trPr.append(th)
    for i, htxt in enumerate(headers):
        c = hdr.cells[i]
        c.width = widths[i]
        _shd(c._tc.get_or_add_tcPr(), header_fill)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run(p, htxt, size=header_size, bold=True,
            color=RGBColor(0xFF, 0xFF, 0xFF), font=F_TIT)
    for ri, rowdata in enumerate(rows):
        r = t.add_row()
        for i, val in enumerate(rowdata):
            c = r.cells[i]
            c.width = widths[i]
            if ri % 2 == 1:
                _shd(c._tc.get_or_add_tcPr(), "F5F5F5")
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            for tt in _partes(val):
                txt, bold = tt[0], tt[1]
                it2 = tt[2] if len(tt) > 2 else False
                run(p, txt, size=size, bold=bold, italic=it2)
    if fonte:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(8)
        run(p, fonte, size=9, italic=True, color=CINZA)
    else:
        espaco()
    return t


def check(items, size=10.5):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run(p, "\u2610  ", size=12, bold=True, color=VINHO)
        for tt in _partes(it):
            txt, bold = tt[0], tt[1]
            it2 = tt[2] if len(tt) > 2 else False
            run(p, txt, size=size, bold=bold, italic=it2)


def entrega(texto):
    """Faixa que indica em qual avaliacao a etapa e cobrada."""
    t, c = _one_cell_box("8D0333", "8D0333")
    p = c.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run(p, "ENTREGA:  ", size=9, bold=True, color=RGBColor(0xD4, 0xB2, 0x77), font=F_TIT)
    run(p, texto, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), font=F_TIT)
    espaco()


# ---------------------------------------------------------------- cabecalho e rodape
hdr = sec.header
hdr.is_linked_to_previous = False
htab = hdr.add_table(rows=1, cols=2, width=CONTENT_W)
htab.autofit = False
htab.columns[0].width = Cm(13.3)
htab.columns[1].width = Cm(2.7)
hc0, hc1 = htab.cell(0, 0), htab.cell(0, 1)
hc0.width, hc1.width = Cm(13.3), Cm(2.7)
hp = hc0.paragraphs[0]
hp.paragraph_format.space_after = Pt(0)
hp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
run(hp, "Ministério da Educação", size=8, font=F_TIT, color=CINZA)
hp2 = hc0.add_paragraph()
hp2.paragraph_format.space_after = Pt(0)
run(hp2, "UNIVERSIDADE FEDERAL DO MARANHÃO", size=11, bold=True, font=F_TIT, color=VINHO)
hp3 = hc0.add_paragraph()
hp3.paragraph_format.space_after = Pt(0)
run(hp3, "CENTRO DE CIÊNCIAS SOCIAIS · CURSO DE ADMINISTRAÇÃO (BACHARELADO)",
    size=7.5, font=F_TIT, color=CINZA)
hp4 = hc0.add_paragraph()
hp4.paragraph_format.space_after = Pt(0)
run(hp4, "MÉTODOS E TÉCNICAS DE PESQUISA QUANTITATIVA", size=7.5, font=F_TIT, color=CINZA)
hpl = hc1.paragraphs[0]
hpl.paragraph_format.space_after = Pt(0)
hpl.alignment = WD_ALIGN_PARAGRAPH.RIGHT
if os.path.exists(LOGO):
    hpl.add_run().add_picture(LOGO, height=Cm(1.9))
hsep = hdr.add_paragraph()
hsep.paragraph_format.space_before = Pt(2)
hsep.paragraph_format.space_after = Pt(0)
run(hsep, "", size=1)
_pborder(hsep, "8D0333", 8)

ftr = sec.footer
ftr.is_linked_to_previous = False
fp = ftr.paragraphs[0]
fp.paragraph_format.space_before = Pt(0)
fp.paragraph_format.space_after = Pt(0)
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp, "Template de projeto de pesquisa quantitativa  ·  Prof. Dr. Tadeu Gomes Teixeira  "
        "·  UFMA  ·  SEMPRE+ QUALIDADE INOVAÇÃO INCLUSÃO  ·  pág. ",
    size=7.5, color=CINZA, font=F_TIT)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
fr = OxmlElement("w:r")
frPr = OxmlElement("w:rPr")
sz = OxmlElement("w:sz")
sz.set(qn("w:val"), "15")
frPr.append(sz)
fr.append(frPr)
ft = OxmlElement("w:t")
ft.text = "1"
fr.append(ft)
fld.append(fr)
fp._p.append(fld)

# ================================================================ CAPA
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(30)
p.paragraph_format.space_after = Pt(0)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run(p, "TEMPLATE  ·  PREENCHIMENTO GUIADO", size=10, bold=True, color=DOURADO, font=F_TIT)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(0)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run(p, "Projeto de pesquisa", size=30, bold=True, color=VINHO, font=F_TIT)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(10)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run(p, "quantitativa", size=30, bold=True, color=VINHO, font=F_TIT)
_pborder(p, "D4B277", 12)

para("Da escolha do tema à apresentação dos resultados, etapa por etapa, com o que entra em "
     "cada campo, um exemplo preenchido e a conferência de cada passo.",
     size=12, color=CINZA, after=18)

campo("IDENTIFICAÇÃO", linhas=1, dica="nome completo")
campo(None, linhas=1, dica=None)

tabela(
    ["Campo", "Preencha aqui"],
    [["Nome completo", ""],
     ["Matrícula", ""],
     ["Curso e período", ""],
     ["Semestre", ""],
     ["Título provisório do projeto", ""],
     ["Tema escolhido", ""],
     ["Base de dados principal", ""],
     ["Data da 1ª versão", ""],
     ["Data da versão final", ""],
     ["Link do notebook no Colab", ""]],
    [5, 11], legenda="Quadro de identificação")

para("Marque abaixo em que ponto do semestre você está:", size=10.5, after=4)
check([[("1ª versão", True), (" — encontro 5 (AV1): etapas 0 a 5 e 7 preenchidas", False)],
       [("2ª versão", True), (" — encontro 9 (AV2): etapas 6 a 10 preenchidas e as anteriores reescritas", False)],
       [("Versão final", True), (" — encontro 14 (AV3): o documento inteiro, com resultados", False)]])

# ================================================================ COMO USAR
h1("", "Como usar este template")

para("Este documento é o seu projeto em construção. Ele não é um resumo dos slides: é o "
     "formulário que transforma o que foi visto em aula em um projeto seu, campo por campo. "
     "Ele é deliberadamente redundante com os slides — você não deve precisar procurar nada "
     "fora daqui para preencher.")

h2("", "A regra das três caixas")

caixa_info("O que entra aqui.", [
    "Toda seção começa com uma caixa como esta, em vinho. Ela diz, em linguagem direta, o que "
    "aquele campo pede, o que não é aceito e o erro mais comum. Leia antes de escrever."])

caixa_exemplo([
    "A segunda caixa, em dourado, traz um exemplo completo, tirado das bases da disciplina. "
    "Ele serve de régua: se o seu campo está muito mais curto ou muito mais vago que o "
    "exemplo, ainda não está pronto. Não copie o exemplo — ele é de outro projeto."])

campo("A TERCEIRA CAIXA É SUA", linhas=2, dica="escreva aqui")

para("As caixas de borda cinza são para você escrever. Escreva em frases completas, não em "
     "tópicos soltos: o relatório final é feito de frases, e é mais fácil reaproveitar uma "
     "frase já escrita do que transformar um tópico em texto na véspera da entrega.")

h2("", "As três regras do semestre, que valem também aqui")

numbers([
    [("A IA explica, o código calcula e você confere e escreve.", True),
     (" Nenhum número vai para este documento sem ter saído de uma célula executada por você. "
      "Todo uso de assistente é registrado na etapa 16.", False)],
    [("Nenhum número sem a frase.", True),
     (" Um número solto não vale metade. A frase completa passa pelo teste das quatro "
      "perguntas da etapa 12: quanto, em quem, com que incerteza e o que não autoriza.", False)],
    [("Nenhum verbo causal sobre dado observacional.", True),
     (" As palavras “causa”, “provoca”, “faz aumentar” estão proibidas quando ninguém sorteou "
      "quem recebeu o quê. Use “está associado a”, “é maior em”, “acompanha”.", False)],
])

h2("", "O calendário das entregas")

tabela(
    ["Encontro", "Avaliação", "O que deste documento é cobrado", "Peso"],
    [["5", "AV1",
      "Etapas 0 a 5 e 7: tema, problema, justificativa, objetivos, hipóteses, variáveis e a "
      "base escolhida com a célula de carga funcionando. Mais a matriz de amarração (etapa 10) "
      "em versão preliminar", "1/3"],
     ["9", "AV2",
      "Etapas 6 a 10 reescritas: delineamento, fonte conferida contra os seis critérios, "
      "população e amostra com o tamanho calculado, instrumento (se houver) e a matriz "
      "de amarração fechada", "1/3"],
     ["14", "AV3",
      "O documento inteiro, com as etapas 11 a 17: plano de análise executado, as quatro peças "
      "de resultado, discussão, limitações, registro de IA e a apresentação de 3 a 5 minutos",
      "1/3"],
     ["15", "Prova final",
      "Não se entrega projeto. Mas a prova cobra as decisões que você tomou aqui: a escolha da "
      "medida, a escolha do teste e a leitura do resultado", "—"]],
    [2, 2, 10, 1.6], legenda="Quadro 1 — Onde cada etapa é cobrada")

caixa_info("Atenção ao prazo útil.", [
    "A etapa 7 (fonte de dados) é a que mais derruba projeto, e sempre pelo mesmo motivo: a "
    "base tem a variável desejada, mas para outro período ou outro território. Descobrir isso "
    "no encontro 4 custa uma tarde; descobrir no encontro 13 custa o projeto. Preencha a "
    "etapa 7 antes de se apaixonar pela pergunta."])

h2("", "Antes de começar, três coisas")

bullets([
    [("Salve uma cópia deste arquivo com o seu nome", True),
     (" (por exemplo, "
      "“Projeto_MariaSilva.docx”) e mantenha uma versão por entrega. Não apague as anteriores: "
      "a devolutiva do professor é feita sobre elas.", False)],
    [("Abra o notebook do encontro ao lado.", True),
     (" Quase todo campo numérico deste documento tem "
      "uma célula correspondente no Colab. O documento guarda a frase; o notebook guarda a "
      "conta.", False)],
    [("Não deixe campo em branco.", True),
     (" Quando algo não se aplica ao seu projeto — e isso acontece, "
      "sobretudo nas etapas 8 e 9 —, escreva “não se aplica” e diga por quê em uma linha. "
      "Campo vazio é lido como esquecimento; “não se aplica” com motivo é lido como decisão.",
      False)],
])

# ================================================================ ETAPA 0
h1("Etapa 0", "A escolha do tema")
entrega("AV1 — encontro 5. Rascunho discutido na oficina do encontro 3.")

caixa_info("O que entra aqui.", [
    "Um tema, escolhido do cardápio da disciplina ou proposto por você. Tema não é problema: "
    "“sobrevivência de empresas” é tema; “empresas de maior porte sobrevivem mais?” já é "
    "problema. Nesta etapa basta o tema, mais a razão de ele ser pesquisável com os meios "
    "da disciplina."])

h2("0.1", "O cardápio de temas")

para("Qualquer linha abaixo é um tema aprovado, com base pública e acessível pelos notebooks "
     "da disciplina. Você pode propor outro, desde que a base seja pública e carregue por API "
     "ou arquivo aberto — e desde que você mostre a célula de carga funcionando.", after=6)

tabela(
    ["Tema", "Fonte principal", "Pergunta de exemplo"],
    [["Sobrevivência de empresas por porte ou setor", "Demografia das Empresas (SIDRA 9949)",
      "Empresas empregadoras de maior porte sobrevivem mais que as de menor porte?"],
     ["Estrutura empresarial do Maranhão vs. Brasil", "CEMPRE (SIDRA 9582)",
      "A economia maranhense é mais concentrada em comércio?"],
     ["Preço de combustíveis por município", "ANP — SHPC (dados abertos)",
      "O preço da gasolina difere entre os municípios do Maranhão?"],
     ["Inflação e inadimplência", "SGS/BCB (433, 21086)",
      "A inadimplência PJ acompanha a inflação?"],
     ["Crédito e inadimplência empresarial", "SGS/BCB (20543, 21086)",
      "A inadimplência PJ sobe quando o crédito encarece?"],
     ["Universo de empresas por município e atividade", "Receita Federal (CNPJ) e CEMPRE",
      "Quantas empresas do comércio há em São Luís, e como isso se compara ao estado?"],
     ["Setor de serviços: emprego e receita", "PAS (SIDRA 2325–2330)",
      "Que segmento de serviços mais emprega por real de receita?"],
     ["Rentabilidade de companhias abertas por setor", "CVM (dados abertos DFP)",
      "As margens diferem sistematicamente entre setores?"],
     ["Endividamento e desempenho de empresas listadas", "CVM (dados abertos DFP)",
      "Empresas mais endividadas são menos rentáveis?"]],
    [6, 4.5, 6.5], legenda="Quadro 2 — O cardápio de temas da disciplina")

h2("0.2", "O meu tema")

campo("TEMA ESCOLHIDO", linhas=2, dica="uma linha, em substantivos: “o que”, não “se”")

campo("POR QUE ESTE TEMA ME INTERESSA", linhas=3,
      dica="duas ou três frases; vale motivo pessoal, profissional ou de curiosidade")

campo("A QUEM O RESULTADO SERVIRIA", linhas=3,
      dica="um gestor, um órgão, uma empresa, uma política — e que decisão ele tomaria melhor")

h2("0.3", "Três perguntas de triagem")

para("Responda honestamente. Um “não” aqui não reprova o tema, mas indica onde ele precisa "
     "de trabalho antes da etapa 1.", after=6)

tabela(
    ["Pergunta de triagem", "Sim / Não", "Se “não”, o que falta"],
    [["Existe uma base pública que já contém a variável central deste tema?", "", ""],
     ["Essa base cobre o território e o período que me interessam?", "", ""],
     ["Consigo dizer que tabela ou gráfico responderia à minha curiosidade?", "", ""]],
    [8, 2.2, 5.8])

# ================================================================ ETAPA 1
h1("Etapa 1", "Do tema ao problema de pesquisa")
entrega("AV1 — encontro 5. Este é o campo mais importante do projeto.")

caixa_info("O que entra aqui.", [
    "O problema, enunciado como pergunta, mais a prova de que ele passa pelas seis regras de "
    "Gil (2022, seção 2.3). Problema não é tema, não é opinião e não é juízo de valor. Se a "
    "sua pergunta começa com “é melhor” ou “deveria”, ela não é um problema científico — é "
    "uma questão normativa, e nenhuma base de dados a responde."])

h2("1.1", "O funil: do tema ao problema")

para("Preencha as quatro linhas na ordem. Cada uma estreita a anterior. A última é o seu "
     "problema.", after=6)

tabela(
    ["Degrau", "O que é", "No exemplo da aula", "O seu"],
    [["1. Tema", "A área, em substantivos", "Sobrevivência de empresas", ""],
     ["2. Tema delimitado", "Com população, território e período",
      "Empresas empregadoras do comércio, Maranhão, coortes a partir de 2018", ""],
     ["3. Pergunta ampla", "Já é pergunta, mas ainda vaga",
      "O que explica a sobrevivência das empresas?", ""],
     ["4. Problema", "Pergunta precisa, com variáveis mensuráveis",
      "Empresas empregadoras de maior porte sobrevivem mais, em três anos, que as de menor "
      "porte?", ""]],
    [2.5, 3.5, 5.5, 4.5], legenda="Quadro 3 — O funil do problema")

h2("1.2", "A delimitação, em quatro dimensões")

caixa_info("O que entra aqui.", [
    "A regra 5 de Gil exige o problema delimitado a uma dimensão viável. Delimitar é dizer "
    "quatro coisas, e nenhuma delas pode ficar implícita. Escreva cada uma em linguagem de "
    "dados: “empresas com CNPJ ativo”, e não “empresas de verdade”."])

tabela(
    ["Dimensão", "A pergunta que ela responde", "No exemplo", "O seu"],
    [["População", "Quem são os casos? Qual é a unidade de observação?",
      "Empresas empregadoras do comércio (seção G da CNAE)", ""],
     ["Território", "Que recorte geográfico?", "Maranhão", ""],
     ["Período", "Que anos ou meses, exatamente?", "Coortes de 2018 a 2021", ""],
     ["Aspecto", "Que faceta do fenômeno, e não o fenômeno todo?",
      "Sobrevivência em três anos, e não desempenho em geral", ""]],
    [2.5, 4.5, 4.5, 4.5], legenda="Quadro 4 — A delimitação em quatro dimensões")

h2("1.3", "O meu problema de pesquisa")

campo("PROBLEMA, ENUNCIADO COMO PERGUNTA", linhas=3,
      dica="uma frase, terminada em sinal de interrogação")

caixa_exemplo([
    "“Empresas empregadoras de maior porte do comércio maranhense sobrevivem mais, ao fim de "
    "três anos, que as de menor porte?”"])

h2("1.4", "As seis regras de Gil, uma a uma")

para("Para cada regra, marque a caixa e escreva, na coluna da direita, a evidência de que o "
     "seu problema a cumpre. A caixa marcada sem justificativa não conta.", after=6)

tabela(
    ["#", "Regra (Gil, 2022, seção 2.3)", "O que significa na prática",
     "☐", "Como o meu problema cumpre"],
    [["1", "Formulado como pergunta",
      "“Vou pesquisar o desemprego” esclarece pouco; “em que setores o desemprego é mais "
      "alto?” já é problema", "", ""],
     ["2", "Claro e preciso",
      "Termos do cotidiano são ambíguos (“organização”, “desempenho”): defina-os "
      "operacionalmente, dizendo como serão medidos", "", ""],
     ["3", "Com base empírica",
      "Não pede juízo de valor (“é melhor?”, “deve?”); refere-se a fatos observáveis", "", ""],
     ["4", "Suscetível de solução",
      "Existe técnica e dado para respondê-lo agora, não numa tecnologia futura", "", ""],
     ["5", "Delimitado a uma dimensão viável",
      "População, território, período e aspecto definidos conforme os meios disponíveis", "",
      ""],
     ["6", "Ético",
      "Respeito aos participantes, riscos contra benefícios, relevância social "
      "(Res. CNS 466/2012)", "", ""]],
    [0.7, 3.2, 6.0, 0.8, 5.3], legenda="Quadro 5 — Conferência pelas seis regras")

caixa_info("O teste rápido das regras 4 e 5.", [
    "Consegue dizer ", ("que tabela ou que base responderia a sua pergunta", True),
    ", com o nome da fonte e o período? Se não faz ideia, o problema ainda não está pronto — "
    "e nenhuma reescrita de frase resolve isso. Volte à etapa 0.3."])

campo("A TABELA OU BASE QUE RESPONDE À MINHA PERGUNTA", linhas=2,
      dica="nome da fonte, tabela ou código, e período disponível")

h2("1.5", "As definições operacionais")

caixa_info("O que entra aqui.", [
    "A regra 2 em detalhe. Definição operacional é a que indica ", ("como o fenômeno é medido",
                                                                    True),
    ". “Empresa inovadora” é ambíguo até virar “declarou à PINTEC produto novo no triênio”. "
    "Liste todos os termos da sua pergunta que precisam de definição, e não só os que parecem "
    "difíceis: “porte”, “desempenho” e “sobrevivência” parecem óbvios e não são."])

tabela(
    ["Termo da minha pergunta", "Definição operacional", "Como é medido",
     "Coluna ou série na base"],
    [["Porte (exemplo)", "Faixa de pessoal ocupado assalariado declarada no CEMPRE",
      "Faixa ordinal de 0 a 4, de 5 a 9, …, 500 ou mais",
      "faixa_pessoal, SIDRA 9949"],
     ["", "", "", ""], ["", "", "", ""], ["", "", "", ""], ["", "", "", ""]],
    [3.2, 4.6, 4.2, 4.0], legenda="Quadro 6 — Definições operacionais")

h2("1.6", "Problemas bons e ruins, lado a lado")

para("Use esta tabela como espelho. Se o seu problema se parece com a coluna da esquerda, "
     "reescreva-o antes de seguir.", after=6)

tabela(
    ["Problema ruim", "Por que não serve", "Reescrito"],
    [["O crédito é importante para as empresas?",
      "Juízo de valor; não há variável a medir",
      "A inadimplência PJ sobe quando a taxa média de juros do crédito PJ sobe?"],
     ["Como melhorar a gestão das empresas maranhenses?",
      "Verbo de intervenção; pede ação, não investigação",
      "A margem líquida difere entre os setores das companhias abertas?"],
     ["O que explica o preço dos combustíveis?",
      "Amplo demais; nenhuma base responde por inteiro",
      "O preço médio da gasolina difere entre a Ilha de São Luís e o interior do Maranhão?"],
     ["Empresas inovadoras são melhores?",
      "Dois termos ambíguos e nenhuma definição operacional",
      "Empresas que declararam produto novo à PINTEC têm receita média maior?"]],
    [4.2, 4.2, 7.6], legenda="Quadro 7 — Problemas ruins e a reescrita de cada um")

# ================================================================ ETAPA 2
h1("Etapa 2", "Justificativa e literatura")
entrega("AV1 — encontro 5.")

caixa_info("O que entra aqui.", [
    "Por que a pergunta merece ser respondida, e o que já se sabe sobre ela. Não é preciso "
    "revisão sistemática: três referências bem lidas bastam, desde que você diga o que cada "
    "uma já respondeu e o que ela deixou em aberto. Justificativa não é elogio ao tema — é o "
    "argumento de que a resposta muda alguma decisão."])

h2("2.1", "A relevância, em três frases")

campo("RELEVÂNCIA PRÁTICA", linhas=3,
      dica="que decisão de gestão ou de política fica melhor informada com a resposta")

campo("RELEVÂNCIA ACADÊMICA", linhas=3,
      dica="que lacuna, divergência ou falta de evidência local a pesquisa ocupa")

campo("VIABILIDADE", linhas=3,
      dica="por que você consegue fazer isto em um semestre, com esta base e estas ferramentas")

h2("2.2", "As referências que sustentam o problema")

tabela(
    ["Autor (ano)", "O que a obra responde", "O achado principal, em uma frase",
     "Como eu uso no projeto"],
    [["", "", "", ""], ["", "", "", ""], ["", "", "", ""]],
    [2.8, 4.0, 5.2, 4.0], legenda="Quadro 8 — As três referências mínimas",
    fonte="Transcreva cada uma em ABNT na etapa 15. Aqui basta autor e ano.")

h2("2.3", "A lacuna")

campo("O QUE A LITERATURA JÁ RESPONDE", linhas=3)
campo("O QUE AINDA FALTA, E QUE A MINHA PESQUISA FAZ", linhas=3,
      dica="pode ser o recorte: ninguém fez isso para o Maranhão, ou para este período")

# ================================================================ ETAPA 3
h1("Etapa 3", "Objetivos")
entrega("AV1 — encontro 5.")

caixa_info("O que entra aqui.", [
    "Um objetivo geral, que é a sua pergunta transformada em ação de pesquisa, e de dois a "
    "quatro objetivos específicos, que são os passos para chegar lá. Cada objetivo começa por "
    "um verbo no infinitivo, e o verbo escolhido já diz se o objetivo é de pesquisa ou não."])

h2("3.1", "Os verbos denunciam")

tabela(
    ["Verbos de pesquisa ✓", "Verbos vagos ✗", "Verbos de intervenção ✗"],
    [["descrever, identificar, comparar, verificar, avaliar, analisar, mensurar, estimar, "
      "testar, relacionar",
      "pesquisar, entender, conhecer, estudar, abordar, refletir sobre",
      "conscientizar, melhorar, propor soluções, implantar, otimizar, capacitar"]],
    [5.5, 5.0, 5.5], legenda="Quadro 9 — Os três grupos de verbos")

para("Gil recomenda verbos que não possibilitam muitas interpretações. Verbo vago deixa o "
     "leitor sem saber o que será entregue; verbo de intervenção denuncia que se planeja uma "
     "ação, e não uma investigação. A pesquisa pode fundamentar a ação, mas primeiro precisa "
     "responder à pergunta.")

h2("3.2", "O meu objetivo geral")

campo("OBJETIVO GERAL", linhas=2,
      dica="um verbo da coluna ✓ + o que + em quem + onde + quando")

caixa_exemplo([
    "“Comparar as taxas de sobrevivência em três anos entre empresas empregadoras de maior e "
    "de menor porte no comércio maranhense, nas coortes de 2018 a 2021.”"])

h2("3.3", "Os meus objetivos específicos")

caixa_info("O que entra aqui.", [
    "De dois a quatro. Cada um é um passo executável, e juntos eles esgotam o geral — nada "
    "sobra, nada falta. Na prática, o primeiro costuma ser descritivo (etapa 11.1), o segundo "
    "inferencial (etapa 11.3) e o terceiro de modelagem (etapa 11.4). Isso não é regra, mas é "
    "um bom esqueleto."])

campo("OBJETIVO ESPECÍFICO 1", linhas=2)
campo("OBJETIVO ESPECÍFICO 2", linhas=2)
campo("OBJETIVO ESPECÍFICO 3", linhas=2)
campo("OBJETIVO ESPECÍFICO 4 (se houver)", linhas=2)

h2("3.4", "Conferência dos objetivos")

check([
    "Cada objetivo começa por um verbo da coluna ✓ do Quadro 9.",
    "O objetivo geral responde exatamente à pergunta da etapa 1.3 — nem mais, nem menos.",
    "Juntos, os específicos cumprem o geral; nenhum deles vai além do que a base permite.",
    "Nenhum objetivo exige dado que eu não tenho (confira contra a etapa 7).",
    "Nenhum objetivo propõe intervenção, melhoria ou solução.",
])

# ================================================================ ETAPA 4
h1("Etapa 4", "Hipóteses")
entrega("AV1 — encontro 5, reescrita na AV2.")

caixa_info("O que entra aqui.", [
    "A resposta provável à sua pergunta, enunciada antes de olhar o resultado. Uma hipótese "
    "é uma ", ("afirmação sobre a relação entre duas variáveis", True),
    ", com direção declarada. Não toda pesquisa precisa de hipótese: um projeto puramente "
    "descritivo pode não ter. Mas se o seu projeto tem comparação de grupos ou associação, "
    "ele tem hipótese — mesmo que você não a escreva, e aí o problema é que o teste vai "
    "respondê-la sem que você saiba qual era."])

h2("4.1", "As minhas hipóteses")

tabela(
    ["#", "Enunciado da hipótese", "Variável independente", "Variável dependente",
     "Direção esperada"],
    [["H1 (ex.)", "Quanto maior a faixa de pessoal, maior a taxa de sobrevivência em 3 anos",
      "Faixa de pessoal (ordinal)", "Taxa de sobrevivência (razão)", "Positiva"],
     ["H1", "", "", "", ""], ["H2", "", "", "", ""], ["H3", "", "", "", ""]],
    [1.2, 5.2, 3.2, 3.2, 2.2], legenda="Quadro 10 — As hipóteses do projeto")

h2("4.2", "Os seis requisitos da hipótese aplicável")

para("Gil (2022, seção 3.4). Confira cada hipótese contra os seis. A hipótese que falha em "
     "qualquer um não é testável com os meios da disciplina.", after=6)

check([
    [("Conceitualmente clara.", True),
     (" Os termos têm definição operacional (etapa 1.5), e não só dicionário.", False)],
    [("Referida a fatos empíricos.", True),
     (" Nada de juízo de valor: “as empresas deveriam” não é hipótese.", False)],
    [("Específica.", True),
     (" Diz quanto, em quem e em que direção, e não apenas “há relação”.", False)],
    [("Relacionada a técnicas disponíveis.", True),
     (" Existe um teste, visto nos encontros 10 a 13, que a decide.", False)],
    [("Relacionada a uma teoria.", True),
     (" Vem de uma leitura, de uma observação ou de uma experiência, e não do nada.", False)],
    [("Verificável com os dados que eu tenho.", True),
     (" As duas variáveis existem na base da etapa 7, no mesmo recorte.", False)],
])

h2("4.3", "Três defeitos clássicos, e como evitá-los")

tabela(
    ["Defeito", "Exemplo do defeito", "Corrigido"],
    [["Tautologia: a hipótese é verdadeira por definição",
      "“Empresas que sobrevivem têm mais chance de continuar abertas”",
      "“Empresas de maior porte sobrevivem mais em três anos”"],
     ["Impossível de refutar: nenhum dado a contraria",
      "“O porte influencia de alguma forma a sobrevivência”",
      "“A taxa de sobrevivência cresce com a faixa de pessoal”"],
     ["Juízo de valor disfarçado",
      "“Empresas pequenas são mal geridas”",
      "“Empresas de menor porte têm taxa de sobrevivência menor”"]],
    [4.5, 5.8, 5.7], legenda="Quadro 11 — Defeitos de hipótese")

h2("4.4", "Se o meu projeto não tem hipótese")

campo("MEU PROJETO É DESCRITIVO E NÃO TEM HIPÓTESE PORQUE…", linhas=3,
      dica="preencha só se for o caso; diga qual pergunta descritiva substitui a hipótese")

h2("4.5", "Da hipótese à hipótese estatística")

caixa_info("Isto você preenche depois do encontro 10.", [
    "A hipótese da etapa 4.1 é a sua. A hipótese estatística é a forma em que o teste a "
    "escreve, e ela sempre vem em par: a nula, que afirma a ausência de diferença ou de "
    "associação, e a alternativa, que é a sua. O teste nunca prova a sua; ele só decide se "
    "rejeita a nula."])

tabela(
    ["Hipótese", "Hipótese nula (H₀)", "Hipótese alternativa (H₁)", "Nível de significância"],
    [["H1 (ex.)", "As taxas médias de sobrevivência são iguais entre as faixas de porte",
      "As taxas médias diferem entre as faixas de porte", "α = 0,05"],
     ["H1", "", "", ""], ["H2", "", "", ""], ["H3", "", "", ""]],
    [1.5, 5.2, 5.2, 3.1], legenda="Quadro 12 — As hipóteses estatísticas")

# ================================================================ ETAPA 5
h1("Etapa 5", "Variáveis e níveis de mensuração")
entrega("AV1 — encontro 5.")

caixa_info("O que entra aqui.", [
    "Todas as variáveis do projeto, com o nível de mensuração de cada uma. Esta etapa é o "
    "gargalo técnico do projeto inteiro, porque ", ("o nível decide o que se pode calcular",
                                                    True),
    ". Média de variável nominal é erro; mediana de variável nominal é erro; frequência de "
    "variável de razão contínua, sem faixas, não diz nada. Preencha com a base aberta ao lado."])

h2("5.1", "A regra da medida")

tabela(
    ["Nível", "O que distingue", "Exemplo", "Medida de centro", "O que NÃO fazer"],
    [["Nominal", "Só identifica categorias, sem ordem", "Setor, UF, bandeira do posto",
      "Moda", "Média, mediana, desvio padrão"],
     ["Ordinal", "Ordena, mas o intervalo não é constante",
      "Faixa de pessoal, escala de concordância", "Moda, mediana",
      "Média (com ressalva), desvio padrão"],
     ["Intervalar", "Intervalos constantes, zero arbitrário", "Temperatura, escore z",
      "Média, mediana, moda", "Razões (“duas vezes mais”)"],
     ["Razão", "Intervalos constantes e zero absoluto", "Receita, preço, número de empregados",
      "Média, mediana, moda", "Tratar como categoria sem faixas"]],
    [1.8, 4.0, 3.6, 3.0, 3.6], legenda="Quadro 13 — O nível decide a medida")

h2("5.2", "As variáveis do meu projeto")

tabela(
    ["Variável", "Papel", "Nível", "Unidade", "Coluna na base", "Estatística que vou usar"],
    [["Taxa de sobrevivência (ex.)", "Dependente", "Razão", "%", "taxa_sobrev_3a",
      "Média, desvio padrão, teste t"],
     ["Faixa de pessoal (ex.)", "Independente", "Ordinal", "Faixa", "faixa_pessoal",
      "Moda, mediana, qui-quadrado"],
     ["", "", "", "", "", ""], ["", "", "", "", "", ""], ["", "", "", "", "", ""],
     ["", "", "", "", "", ""], ["", "", "", "", "", ""], ["", "", "", "", "", ""]],
    [3.4, 2.2, 1.6, 1.6, 2.8, 4.4],
    legenda="Quadro 14 — Variáveis, níveis e estatísticas previstas",
    fonte="Papel: dependente (resposta), independente (explicativa) ou de controle. "
          "Nível: nominal, ordinal, intervalar ou de razão.")

h2("5.3", "Conferência das variáveis")

check([
    "Cada variável tem nome, nível, unidade e a coluna exata da base.",
    "A variável dependente é a que está na minha pergunta como “o que se quer explicar”.",
    "Nenhuma estatística da última coluna é proibida para aquele nível (confira o Quadro 13).",
    "Toda variável que aparece numa hipótese (etapa 4.1) está nesta tabela.",
    "Nenhuma variável desta tabela ficou sem uso nos objetivos (etapa 3).",
])

# ================================================================ ETAPA 6
h1("Etapa 6", "Delineamento da pesquisa")
entrega("AV2 — encontro 9.")

caixa_info("O que entra aqui.", [
    "O tipo de pesquisa e a justificativa da escolha. O delineamento é a decisão que amarra "
    "todas as seguintes: ele define a população, que define as fontes, que definem o "
    "instrumento, que define a análise. Errar aqui contamina todas as etapas seguintes."])

h2("6.1", "O meu delineamento")

para("Marque um — e apenas um — como principal.", after=6)

check([
    [("Pesquisa documental ou com dados secundários.", True),
     (" Uso base já existente, produzida por outro (IBGE, BCB, CVM, ANP). É o caminho "
      "padrão da disciplina.", False)],
    [("Levantamento (survey).", True),
     (" Coleto dado primário por questionário. Exige a etapa 9 inteira e o cuidado ético do "
      "encontro 8.", False)],
    [("Pesquisa ex-post-facto.", True),
     (" Comparo grupos que já diferem por algo que não foi manipulado por mim.", False)],
    [("Estudo de caso com dado quantitativo.", True),
     (" Uma organização, muitas medidas. Atenção: não permite generalizar.", False)],
    [("Experimento ou quase-experimento.", True),
     (" Há sorteio ou designação de tratamento. Raro em um semestre; fale com o professor "
      "antes.", False)],
])

campo("POR QUE ESTE DELINEAMENTO, E NÃO OUTRO", linhas=3,
      dica="ligue a escolha à pergunta da etapa 1.3 e aos meios de que você dispõe")

h2("6.2", "Temporalidade")

check([
    [("Transversal.", True), (" Um momento no tempo, muitos casos.", False)],
    [("Longitudinal / série temporal.", True),
     (" O mesmo objeto medido em vários períodos.", False)],
    [("Painel.", True), (" Muitos casos, medidos em vários períodos.", False)],
])

campo("O MOMENTO OU O PERÍODO EXATO", linhas=2,
      dica="“dezembro de 2024”, “2018 a 2024, mensal”, “coortes de 2018 a 2021”")

h2("6.3", "A cadeia de decisões")

para("Preencha a cadeia para conferir se nada foi decidido fora de ordem. Se uma célula não "
     "combina com a vizinha, há um buraco no projeto.", after=6)

tabela(
    ["O problema define…", "que define…", "que define…", "que define…"],
    [["o delineamento", "a população e as fontes", "o instrumento ou a extração",
      "a análise prevista"],
     ["", "", "", ""]],
    [4, 4, 4, 4], legenda="Quadro 15 — As decisões encadeadas")

# ================================================================ ETAPA 7
h1("Etapa 7", "A fonte de dados")
entrega("AV1 (preliminar) e AV2 (completa). A célula de carga tem de funcionar na AV1.")

caixa_info("O que entra aqui.", [
    "A identificação completa da sua base e a conferência dela contra os seis critérios do "
    "encontro 4. Uma fonte não é boa por ser grande: é boa por ", ("responder à pergunta", True),
    ". Esta etapa se preenche com o notebook aberto, e não de memória."])

h2("7.1", "Identificação da base")

tabela(
    ["Campo", "Preencha"],
    [["Nome da base", ""],
     ["Produtor (órgão ou instituição)", ""],
     ["Tabela, série ou código", ""],
     ["Endereço (URL ou endpoint da API)", ""],
     ["Forma de acesso (API, CSV aberto, download manual)", ""],
     ["Período coberto pela base", ""],
     ["Período que eu vou usar", ""],
     ["Granularidade (empresa, município, UF, país)", ""],
     ["Número de observações carregadas (n)", ""],
     ["Número de colunas usadas", ""],
     ["Separador e codificação do arquivo", ""],
     ["Data do meu acesso", ""],
     ["Citação completa, como vai no relatório", ""]],
    [6.5, 9.5], legenda="Quadro 16 — Identificação da fonte")

h2("7.2", "Os seis critérios de escolha da fonte")

tabela(
    ["#", "Critério", "A pergunta que se faz", "A minha resposta"],
    [["1", "Periodicidade", "De quanto em quanto tempo é atualizada? Serve para o meu "
      "período?", ""],
     ["2", "Granularidade", "A observação é a empresa, o município, a UF ou o país?", ""],
     ["3", "Cobertura", "Quais territórios e quais anos existem de fato?", ""],
     ["4", "Metadado", "Quem produz, com que metodologia, e a base é revisada?", ""],
     ["5", "Licença e citação", "O que posso publicar? Como cito a fonte no relatório?", ""],
     ["6", "Acesso", "Há API, há arquivo para baixar, ou é preciso digitar?", ""]],
    [0.7, 3.0, 6.5, 5.8], legenda="Quadro 17 — A fonte contra os seis critérios")

caixa_info("Os dois critérios que mais reprovam projeto.", [
    ("Cobertura", True),
    (": muita base tem a variável desejada, mas para outro período ou outro território. E ",
     False),
    ("metadado", True),
    (": uma base sem metadado é um número sem origem — e um número sem origem não pode entrar "
     "num relatório, porque não se sabe se a série foi recalculada, se houve quebra de "
     "metodologia ou se o número é preliminar.", False)])

h2("7.3", "Qualidade do dado: os quatro defeitos")

para("Procure os quatro na sua base, antes de qualquer análise. Para cada um, conte, declare "
     "e diga o que fez. A regra é simples: conta-se e declara-se; nunca se preenche por conta "
     "própria.", after=6)

tabela(
    ["Defeito", "Como ele aparece", "Quantos casos na minha base", "O que eu fiz"],
    [["Coluna ou célula vazia",
      "A coluna “Valor de Compra” da ANP está vazia nas 12 planilhas; na CVM, 35 de 451 "
      "companhias não têm setor", "", ""],
     ["Escala e unidade misturadas",
      "Na CVM, 444 companhias declaram em MIL e 7 em UNIDADE. Misturar faz a mediana das "
      "segundas parecer 34 vezes maior. É o mais perigoso dos quatro", "", ""],
     ["Número no padrão nacional",
      "O “6,29” da ANP e a data DD/MM/AAAA: sem decimal=\",\" a coluna chega como texto e a "
      "conta falha ou sai errada", "", ""],
     ["Codificação",
      "A CVM publica em latin-1. Lida com a codificação errada, não dá erro: dá resultado "
      "errado, com acentos quebrados", "", ""]],
    [3.0, 6.5, 3.0, 3.5], legenda="Quadro 18 — Os quatro defeitos de qualidade do dado")

h2("7.4", "Ética dos dados")

caixa_info("O que entra aqui.", [
    "Dado secundário também tem ética. O risco principal é a ", ("reidentificação", True),
    ": cruzar duas bases agregadas pode individualizar uma empresa ou uma pessoa — e isso "
    "basta para violar o sigilo mesmo que nenhuma base, isolada, o violasse. Se o seu dado é "
    "primário, o bloco de baixo é obrigatório."])

check([
    "A base que eu uso é pública e o seu uso para pesquisa é permitido pela licença.",
    "Não publico nenhuma célula com menos casos do que o mínimo que a fonte divulga.",
    "Não cruzo bases de modo que uma empresa ou pessoa específica fique identificável.",
    "Cito a fonte na forma que ela exige, em toda tabela e em toda figura.",
    [("Coleta primária:", True),
     (" o instrumento informa objetivo, uso dos dados, voluntariedade e contato do "
      "responsável.", False)],
    [("Coleta primária:", True),
     (" há consentimento livre e esclarecido, conforme a Res. CNS 466/2012.", False)],
    [("Coleta primária:", True),
     (" nenhum dado identificável é guardado junto com as respostas.", False)],
])

campo("RISCOS ÉTICOS DO MEU PROJETO E O QUE EU FAÇO A RESPEITO", linhas=3)

h2("7.5", "A carga funcionando")

caixa_info("O que a AV1 cobra.", [
    "Não basta dizer qual é a base: a célula que a carrega tem de executar na frente do "
    "professor, do início ao fim, sem erro. Anote abaixo os dados da célula e o que ela "
    "imprimiu."])

tabela(
    ["Item", "Preencha"],
    [["Link do notebook no Colab", ""],
     ["Nome da célula ou seção que carrega a base", ""],
     ["Linhas e colunas impressas pelo shape", ""],
     ["Primeiro número que eu confiro para saber que carregou certo", ""],
     ["Célula de contingência, se a API falhar (arquivo em dados/)", ""]],
    [7, 9], legenda="Quadro 19 — A carga da base")

# ================================================================ ETAPA 8
h1("Etapa 8", "População, amostra e amostragem")
entrega("AV2 — encontro 9, com o relatório de sorteio.")

caixa_info("O que entra aqui.", [
    "Quem são todos os casos possíveis (a população), quem efetivamente entra na análise (a "
    "amostra, ou o censo) e como se passa de um para o outro. Se você usa uma base oficial "
    "completa, o seu caso costuma ser ", ("censo da base", True),
    " — e isso não dispensa a etapa: dispensa o sorteio, mas exige que você diga que a sua "
    "base é a população inteira e o que isso permite e não permite concluir."])

h2("8.1", "A população e a unidade de observação")

tabela(
    ["Campo", "Preencha", "No exemplo"],
    [["Unidade de observação (um caso é…)", "", "Um posto revendedor"],
     ["População-alvo, por extenso", "", "Postos de gasolina do Maranhão em 2025"],
     ["Tamanho da população (N), se conhecido", "", "Não publicado pela ANP"],
     ["Lista ou cadastro que serve de moldura amostral", "", "Pesquisa SHPC da ANP"],
     ["Quem fica de fora da moldura, e por quê", "", "Postos não pesquisados pela ANP"]],
    [5.5, 5.5, 5.0], legenda="Quadro 20 — População e unidade de observação")

h2("8.2", "Censo ou amostra?")

check([
    [("Censo da base.", True),
     (" Uso todos os casos que a base traz. Vou declarar isso e tratar os números como "
      "descritivos da base, não como estimativas de uma população maior.", False)],
    [("Amostra probabilística.", True),
     (" Há sorteio, e por isso posso calcular margem de erro e intervalo de confiança.",
      False)],
    [("Amostra não probabilística.", True),
     (" Não há sorteio. Vou declarar que não cabe generalizar e que a margem de erro não se "
      "aplica.", False)],
])

h2("8.3", "A técnica de amostragem")

tabela(
    ["Grupo", "Técnica", "Quando cabe", "☐"],
    [["Probabilística", "Aleatória simples",
      "Há lista completa e os casos são homogêneos", ""],
     ["Probabilística", "Sistemática",
      "Há lista ordenada e sem periodicidade que coincida com o intervalo de salto", ""],
     ["Probabilística", "Estratificada",
      "Há subgrupos internamente homogêneos e diferentes entre si (setor, porte, região)", ""],
     ["Probabilística", "Por conglomerados",
      "Os casos vêm agrupados naturalmente (municípios, escolas, filiais)", ""],
     ["Não probabilística", "Por conveniência",
      "Quem estava acessível. Não generaliza", ""],
     ["Não probabilística", "Por cotas",
      "Reproduz proporções conhecidas, mas sem sorteio dentro da cota", ""],
     ["Não probabilística", "Intencional ou por julgamento",
      "Casos escolhidos por serem típicos ou extremos", ""],
     ["Não probabilística", "Bola de neve",
      "Um respondente indica o próximo. Populações de difícil acesso", ""]],
    [3.0, 3.2, 8.0, 0.8], legenda="Quadro 21 — As técnicas de amostragem")

campo("A TÉCNICA QUE EU USO E POR QUÊ", linhas=3,
      dica="se houve sorteio, diga a semente usada e onde está a célula que o fez")

h2("8.4", "O tamanho da amostra")

caixa_info("Como se preenche.", [
    "O tamanho não se adivinha: ele sai do intervalo que você aceita. Declare a margem de erro "
    "e o nível de confiança, rode a célula do encontro 8 e anote o que ela devolveu. ",
    ("Não calcule à mão", True),
    " — e não copie um número de outro trabalho."])

tabela(
    ["Parâmetro", "Símbolo", "Valor que eu declarei", "De onde vem"],
    [["Nível de confiança", "1 − α", "", "Decisão minha; 95% é o padrão"],
     ["Valor crítico", "z ou t", "", "Saída da célula"],
     ["Margem de erro aceita", "e", "", "Decisão minha, na unidade da variável"],
     ["Proporção esperada (se a variável é proporção)", "p", "", "Estudo anterior, ou 0,5"],
     ["Desvio padrão esperado (se é média)", "s", "", "Base-piloto ou literatura"],
     ["Tamanho da população, se finita", "N", "", "Quadro 20"],
     ["Tamanho da amostra calculado", "n", "", "Saída da célula do encontro 8"],
     ["Tamanho efetivamente obtido", "n obtido", "", "Contagem no notebook"],
     ["Perda ou não resposta", "—", "", "n calculado menos n obtido, e o motivo"]],
    [6.0, 2.0, 3.5, 4.5], legenda="Quadro 22 — O tamanho da amostra, parâmetro por parâmetro")

campo("A FRASE DO TAMANHO DA AMOSTRA", linhas=3,
      dica="“Para uma margem de ____ e confiança de 95%, a célula devolveu n = ____; obtive ____.”")

h2("8.5", "O que o meu desenho permite dizer")

campo("GENERALIZAÇÃO POSSÍVEL", linhas=3,
      dica="a quem vale o resultado: à base, à população sorteada, ou a ninguém além dos casos observados")

# ================================================================ ETAPA 9
h1("Etapa 9", "Instrumento de coleta (só para dado primário)")
entrega("AV2 — encontro 9.")

caixa_info("Esta etapa pode não se aplicar.", [
    "Se o seu projeto usa base secundária — o caminho padrão da disciplina —, marque “não se "
    "aplica” em 9.1, escreva em uma linha por quê, e passe para a etapa 10. Não deixe em branco."])

h2("9.1", "Aplica-se?")

check([
    [("Não se aplica.", True),
     (" Meu projeto usa dados secundários. A “coleta” é a extração descrita na etapa 7.5.",
      False)],
    [("Aplica-se.", True), (" Eu coleto dado primário, e preencho 9.2 a 9.5.", False)],
])

h2("9.2", "O instrumento")

tabela(
    ["Campo", "Preencha"],
    [["Tipo de instrumento (questionário, formulário, roteiro estruturado)", ""],
     ["Forma de aplicação (presencial, online, telefone)", ""],
     ["Número de itens", ""],
     ["Tempo estimado de resposta", ""],
     ["Escala principal usada", ""],
     ["O instrumento é meu, adaptado ou validado por outro autor?", ""],
     ["Se adaptado ou validado: referência", ""]],
    [8, 8], legenda="Quadro 23 — O instrumento")

h2("9.3", "As escalas")

tabela(
    ["Escala", "O que mede", "Nível de mensuração", "Cuidado principal"],
    [["Nominal (sim/não, categorias)", "Pertencimento a uma categoria", "Nominal",
      "Categorias exaustivas e mutuamente exclusivas"],
     ["Likert (1 a 5, 1 a 7)", "Grau de concordância", "Ordinal",
      "Número de pontos ímpar dá ponto neutro; declare a decisão"],
     ["Diferencial semântico", "Posição entre dois extremos", "Ordinal",
      "Os extremos precisam ser realmente opostos"],
     ["Numérica aberta", "Quantidade, valor, contagem", "Razão",
      "Unidade declarada no enunciado do item"]],
    [3.6, 3.6, 3.0, 5.8], legenda="Quadro 24 — As escalas e o nível que cada uma produz")

h2("9.4", "Matriz item × objetivo × variável")

tabela(
    ["Item nº", "Enunciado do item", "Objetivo específico que ele serve", "Variável", "Escala"],
    [["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""],
     ["", "", "", "", ""], ["", "", "", "", ""]],
    [1.5, 6.0, 3.5, 3.0, 2.0], legenda="Quadro 25 — Cada item serve a um objetivo",
    fonte="Item que não serve a nenhum objetivo sai do instrumento. Objetivo sem item "
          "correspondente fica sem resposta.")

h2("9.5", "Pré-teste")

tabela(
    ["Campo", "Preencha"],
    [["Com quantas pessoas pré-testei", ""],
     ["Quando", ""],
     ["Itens que geraram dúvida", ""],
     ["O que eu mudei depois do pré-teste", ""]],
    [6, 10], legenda="Quadro 26 — O pré-teste")

# ================================================================ ETAPA 10
h1("Etapa 10", "A matriz de amarração")
entrega("AV1 (preliminar) e AV2 (fechada). É o raio-x do projeto.")

caixa_info("O que entra aqui.", [
    "A tabela de verificação da ", ("coerência interna", True),
    " do projeto. Cada linha liga problema → objetivo específico → hipótese → variáveis (e "
    "níveis) → base → técnica de análise. Se qualquer elo falta ou não combina com o vizinho, "
    "o projeto tem um buraco — e é melhor descobri-lo agora do que na análise."])

h2("10.1", "O exemplo, coluna a coluna")

tabela(
    ["Problema", "Objetivo específico", "Hipótese", "Variáveis (nível)", "Base", "Técnica"],
    [["Empresas empregadoras de maior porte sobrevivem mais que as de menor porte?",
      "Comparar taxas de sobrevivência em 3 anos por faixa de pessoal",
      "Quanto maior a faixa, maior a taxa",
      "Faixa de pessoal (ordinal); taxa de sobrevivência (razão)",
      "Demografia das Empresas, SIDRA 9949",
      "Comparação de grupos"]],
    [3.4, 3.0, 2.2, 3.2, 2.4, 1.8], legenda="Quadro 27 — Matriz de amarração preenchida",
    fonte="Na coluna “técnica”, até o encontro 9 basta o tipo: comparação de grupos, "
          "associação entre variáveis, descrição. Os testes específicos entram depois dos "
          "encontros 11 e 12.")

h2("10.2", "A matriz do meu projeto")

para("Uma linha por objetivo específico. Repita o problema em todas as linhas: é ele que cada "
     "linha precisa servir.", after=6)

tabela(
    ["Problema", "Objetivo específico", "Hipótese", "Variáveis (nível)", "Base", "Técnica"],
    [["", "", "", "", "", ""], ["", "", "", "", "", ""], ["", "", "", "", "", ""],
     ["", "", "", "", "", ""]],
    [3.4, 3.0, 2.2, 3.2, 2.4, 1.8], legenda="Quadro 28 — A minha matriz de amarração")

h2("10.3", "O que a matriz revela")

para("Leia a sua matriz procurando os quatro buracos. Marque o que encontrou e escreva o que "
     "fez para fechá-lo.", after=6)

tabela(
    ["Buraco", "Como ele aparece na matriz", "Encontrei?", "O que eu fiz"],
    [["Objetivo sem hipótese", "Linha com objetivo preenchido e hipótese vazia, num projeto "
      "que compara grupos", "", ""],
     ["Hipótese com variável que a base não tem",
      "A coluna de variáveis cita algo que não está no Quadro 14", "", ""],
     ["Técnica incompatível com o nível de mensuração",
      "Média de variável nominal, teste t com variável ordinal de 3 pontos", "", ""],
     ["Problema que nenhuma linha responde por inteiro",
      "Todas as linhas respondem pedaços, e o problema fica maior que a soma delas", "", ""]],
    [3.6, 6.2, 1.8, 4.4], legenda="Quadro 29 — Os quatro buracos")

# ================================================================ ETAPA 11
h1("Etapa 11", "Plano de análise")
entrega("AV3 — encontro 14, mas comece a preencher a partir do encontro 10.")

caixa_info("O que entra aqui.", [
    "O que você vai calcular, nesta ordem, e por quê. O plano se escreve ", ("antes", True),
    " de rodar a análise: é isso que impede a prática de sair caçando o teste que dá o "
    "resultado desejado. Cada linha deste plano vira uma célula do notebook."])

h2("11.1", "Descritiva")

tabela(
    ["Variável", "Nível", "Medida de centro", "Medida de dispersão", "Forma", "Por que esta"],
    [["", "", "", "", "", ""], ["", "", "", "", "", ""], ["", "", "", "", "", ""],
     ["", "", "", "", "", ""], ["", "", "", "", "", ""]],
    [3.2, 1.6, 2.8, 2.8, 2.0, 3.6], legenda="Quadro 30 — O plano da descritiva",
    fonte="Lembrete do encontro 4: em distribuição assimétrica, a média engana; use a mediana "
          "e diga que o fez. Para comparar dispersões de escalas diferentes, use o "
          "coeficiente de variação, não o desvio padrão.")

h2("11.2", "A figura prevista")

tabela(
    ["Campo", "Preencha"],
    [["Tipo de gráfico", ""],
     ["O que vai no eixo x", ""],
     ["O que vai no eixo y", ""],
     ["Por que este tipo, e não outro", ""],
     ["O que a figura deve mostrar, se a hipótese estiver certa", ""]],
    [6, 10], legenda="Quadro 31 — A Figura 1 planejada")

h2("11.3", "Inferência")

caixa_info("Como se escolhe o teste.", [
    "Pelo fluxograma dos encontros 11 e 12, e não pelo que deu significativo. A escolha "
    "depende de três coisas: o nível de mensuração da variável resposta, quantos grupos se "
    "comparam e se os pressupostos do teste paramétrico se sustentam. Quando não se "
    "sustentam, o equivalente de postos entra no lugar — e isso não é derrota, é a escolha "
    "certa."])

tabela(
    ["Comparação", "Resposta", "Grupos", "Teste paramétrico", "Equivalente de postos"],
    [["Uma média contra um valor de referência", "Numérica", "1", "Teste t de uma amostra",
      "Wilcoxon de uma amostra"],
     ["Duas médias independentes", "Numérica", "2", "Teste t de duas amostras (Welch)",
      "Mann-Whitney"],
     ["Duas medidas no mesmo caso", "Numérica", "2 pareados", "Teste t pareado",
      "Wilcoxon pareado"],
     ["Três ou mais médias", "Numérica", "3+", "ANOVA de um fator", "Kruskal-Wallis"],
     ["Duas variáveis categóricas", "Categórica", "—", "Qui-quadrado de independência", "—"],
     ["Distribuição observada contra esperada", "Categórica", "—",
      "Qui-quadrado de bondade de ajuste", "—"],
     ["Associação entre duas numéricas", "Numérica", "—", "Correlação de Pearson",
      "Correlação de Spearman"]],
    [4.2, 2.0, 1.6, 4.2, 4.0], legenda="Quadro 32 — O fluxograma de escolha do teste")

tabela(
    ["Hipótese", "Teste escolhido", "Pressupostos que eu verifiquei", "α", "Célula do notebook"],
    [["H1", "", "", "", ""], ["H2", "", "", "", ""], ["H3", "", "", "", ""]],
    [1.5, 3.5, 5.5, 1.0, 4.5], legenda="Quadro 33 — Os testes do meu projeto")

h2("11.4", "O modelo")

caixa_info("Quando entra um modelo.", [
    "Quando a pergunta é sobre o efeito de uma variável ", ("mantidas as outras", True),
    ". Regressão linear quando a resposta é numérica; regressão logística quando a resposta é "
    "0/1. O relatório nunca traz a saída bruta do pacote: traz a tabela de cinco colunas da "
    "etapa 12.4."])

tabela(
    ["Campo", "Preencha"],
    [["Tipo de modelo (linear simples, linear múltipla, logística)", ""],
     ["Variável resposta", ""],
     ["Variáveis explicativas", ""],
     ["Variáveis de controle, se houver", ""],
     ["Variável de evento 0/1, se logística", ""],
     ["O que eu espero que o coeficiente principal mostre", ""]],
    [7, 9], legenda="Quadro 34 — O modelo planejado")

h2("11.5", "Conferência do plano")

check([
    "Cada objetivo específico (etapa 3.3) tem pelo menos uma linha de plano correspondente.",
    "Cada teste do Quadro 33 foi escolhido pelo fluxograma do Quadro 32, e eu sei dizer por "
    "qual caminho.",
    "Nenhuma estatística do plano é proibida para o nível da variável (Quadro 13).",
    "Eu declarei o α antes de ver o valor-p.",
    "O plano cabe no notebook: toda linha tem uma célula correspondente, já identificada.",
])

# ================================================================ ETAPA 12
h1("Etapa 12", "Resultados: as quatro peças")
entrega("AV3 — encontro 14. Peso maior da nota do relatório.")

caixa_info("O que entra aqui.", [
    "Os resultados, sempre nestas quatro peças e sempre nesta ordem. Cada peça vem com a sua "
    "frase de leitura, e ", ("cada número sai de uma célula executada", True),
    " do seu notebook. Toda tabela e toda figura têm número, título, unidade, n e fonte — e "
    "são citadas no texto. Uma figura que o texto não menciona não deveria estar no relatório."])

h2("12.1", "O mapa das quatro peças")

tabela(
    ["Peça", "O que responde", "No exemplo da aula"],
    [["Tabela 1", "Quem são os casos: n, valor típico, variação",
      "236 postos; média de R$ 5,97 (Ilha R$ 5,70; interior R$ 6,31)"],
     ["Figura 1", "A forma da distribuição em cada grupo",
      "Boxplot por região: as caixas não se sobrepõem"],
     ["Inferência", "O que a amostra autoriza dizer",
      "Diferença de R$ 0,61; t = −19,44; p < 0,001"],
     ["Modelo", "O efeito de cada variável, mantidas as outras",
      "Marca +R$ 0,05 (p = 0,079); Ilha −R$ 0,61; R² = 0,66"]],
    [2.2, 5.0, 8.8], legenda="Quadro 35 — As quatro peças do resultado")

h2("12.2", "O teste das quatro perguntas")

caixa_info("A regra que decide a nota.", [
    "Toda frase de resultado diz quatro coisas. Se faltar uma, reescreva. ",
    ("Número sem frase vale menos da metade", True),
    ", e verbo causal sobre dado observacional custa os pontos do modelo."])

tabela(
    ["#", "A pergunta", "O que a frase precisa trazer"],
    [["1", "Quanto?", "O número, com a unidade"],
     ["2", "Em quem?", "O n, o recorte e o período"],
     ["3", "Com que incerteza?", "O intervalo de confiança, o erro padrão ou o valor-p"],
     ["4", "O que não autoriza?", "Causa, generalização, o que ficou de fora"]],
    [0.8, 4.0, 11.2], legenda="Quadro 36 — O teste das quatro perguntas")

caixa_exemplo([
    "“O preço médio da gasolina é R$ 0,61 por litro maior no interior que na Ilha de São Luís "
    "(236 postos pesquisados pela ANP no Maranhão em 2025). Os intervalos de 95% não se "
    "sobrepõem (Ilha: R$ 5,68 a R$ 5,72; interior: R$ 6,25 a R$ 6,37), e o teste t de Welch "
    "rejeita a igualdade (t = −19,44; p < 0,001). O dado não separa frete, concorrência e "
    "escala: não é causa.”"])

h2("12.3", "Tabela 1 e Figura 1")

campo("TÍTULO DA TABELA 1", linhas=1,
      dica="“Tabela 1 — …”, com a unidade, o n e o período")

tabela(
    ["Variável ou grupo", "n", "Média", "Mediana", "Desvio padrão", "Mínimo", "Máximo"],
    [["", "", "", "", "", "", ""], ["", "", "", "", "", "", ""], ["", "", "", "", "", "", ""],
     ["", "", "", "", "", "", ""]],
    [4.0, 1.4, 2.2, 2.2, 2.4, 1.9, 1.9], legenda="Quadro 37 — A minha Tabela 1",
    fonte="Fonte: ____________ . Elaboração própria.")

campo("A FRASE DE LEITURA DA TABELA 1", linhas=3,
      dica="passe pelas quatro perguntas do Quadro 36")

campo("TÍTULO DA FIGURA 1 E O QUE ELA MOSTRA", linhas=3,
      dica="descreva em palavras; o gráfico em si fica no notebook e no relatório")

campo("A FRASE DE LEITURA DA FIGURA 1", linhas=3)

h2("12.4", "Inferência e modelo")

tabela(
    ["Hipótese", "Teste aplicado", "Estatística", "Valor-p", "Decisão sobre H₀"],
    [["H1", "", "", "", ""], ["H2", "", "", "", ""], ["H3", "", "", "", ""]],
    [1.5, 4.5, 3.0, 2.5, 4.5], legenda="Quadro 38 — O resultado dos testes")

campo("A FRASE DE LEITURA DA INFERÊNCIA", linhas=3)

para("A tabela do modelo tem sempre estas cinco colunas — nunca a saída bruta do pacote. "
     "A coluna “efeito” é a frase, em palavras, do que aquela linha significa.", after=6)

tabela(
    ["variável", "estimativa", "erro padrão", "valor-p", "efeito, em palavras"],
    [["intercepto (ex.)", "6,2901", "0,0252", "< 0,001", "preço previsto: branca, interior"],
     ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""]],
    [3.2, 2.4, 2.4, 2.0, 6.0], legenda="Quadro 39 — O modelo em cinco colunas",
    fonte="Anote também o R² (ou o pseudo-R², se logística) e o n do modelo: "
          "R² = ______ ; n = ______ .")

campo("A FRASE DE LEITURA DO MODELO", linhas=3,
      dica="uma frase por linha que importa; nenhum verbo causal")

h2("12.5", "Conferência dos resultados")

check([
    "As quatro peças estão presentes, nesta ordem.",
    "Cada tabela e cada figura tem número, título, unidade, n e fonte.",
    "Toda tabela e toda figura é citada no texto.",
    "Cada peça tem a sua frase, e cada frase passa pelo teste das quatro perguntas.",
    "Nenhum número do relatório deixou de sair de uma célula executada por mim.",
    "Nenhuma frase usa “causa”, “provoca”, “faz aumentar” sobre dado observacional.",
])

# ================================================================ ETAPA 13
h1("Etapa 13", "Discussão, limitações e implicação")
entrega("AV3 — encontro 14.")

caixa_info("O que entra aqui.", [
    "A volta à pergunta. A discussão responde, em uma frase, o que a etapa 1.3 perguntou; "
    "depois confronta esse achado com a literatura da etapa 2; depois declara os limites; e "
    "fecha com o que um gestor faria com isso. A pergunta que o professor faz a quase todos "
    "na apresentação é esta: ", ("“o que esse resultado não autoriza dizer?”", True),
    " Tenha a resposta pronta."])

h2("13.1", "A resposta à pergunta")

campo("A MINHA PERGUNTA, RESPONDIDA EM UMA FRASE", linhas=3,
      dica="repita a pergunta da etapa 1.3 e responda-a com o número principal")

h2("13.2", "O diálogo com a literatura")

tabela(
    ["Referência", "O que ela achou", "O que eu achei", "Convergente ou divergente? Por quê?"],
    [["", "", "", ""], ["", "", "", ""], ["", "", "", ""]],
    [3.0, 4.0, 4.0, 5.0], legenda="Quadro 40 — O achado contra a literatura")

h2("13.3", "As limitações")

caixa_info("Como se escreve uma limitação.", [
    "Limitação não é desculpa e não é autodepreciação. É a declaração precisa do que o seu "
    "desenho não alcança. Três tipos sempre cabem: o limite ", ("da base", True),
    " (o que ela não mede), o limite ", ("do desenho", True),
    " (o que ele não permite concluir) e o limite ", ("do recorte", True),
    " (a quem o resultado não se aplica)."])

campo("LIMITAÇÃO 1 — DA BASE", linhas=2,
      dica="o que a base não mede, quantos casos faltaram, que defeito do Quadro 18 sobrou")

campo("LIMITAÇÃO 2 — DO DESENHO", linhas=2,
      dica="observacional, transversal, sem sorteio, variáveis de confusão não medidas")

campo("LIMITAÇÃO 3 — DO RECORTE", linhas=2,
      dica="território, período e população a que o resultado não se estende")

h2("13.4", "O que o resultado não autoriza dizer")

campo("RESPOSTA PRONTA PARA A PERGUNTA DO PROFESSOR", linhas=3,
      dica="três coisas que alguém poderia concluir do seu resultado e que o seu dado não sustenta")

caixa_info("O exemplo da variável de confusão.", [
    "No caso dos postos, a região e a bandeira não foram sorteadas. Se os postos do interior "
    "são, ao mesmo tempo, mais distantes da base de distribuição, o frete entra na conta sem "
    "aparecer no modelo: é uma ", ("variável de confusão", True),
    ", e ela é a razão pela qual associação não é causa."])

h2("13.5", "A implicação para a gestão")

campo("O QUE UM GESTOR FARIA COM ESTE RESULTADO", linhas=3,
      dica="uma decisão concreta, não um conselho genérico")

# ================================================================ ETAPA 14
h1("Etapa 14", "Cronograma")
entrega("AV1 — encontro 5, atualizado nas versões seguintes.")

tabela(
    ["Encontro", "O que eu faço no projeto", "Prazo", "Feito?"],
    [["1 a 2", "Escolho o tema (etapa 0) e leio a base pela primeira vez", "", ""],
     ["3", "Escrevo o problema, os objetivos e as hipóteses (etapas 1, 3 e 4)", "", ""],
     ["4", "Confiro a fonte pelos seis critérios e monto a matriz (etapas 7 e 10)", "", ""],
     ["5", "Entrego a AV1 com a carga funcionando", "", ""],
     ["6 a 7", "Defino o delineamento e a população (etapas 6 e 8.1)", "", ""],
     ["8", "Calculo o tamanho da amostra e fecho o instrumento (etapas 8.4 e 9)", "", ""],
     ["9", "Entrego a AV2 com a matriz fechada", "", ""],
     ["10 a 11", "Escolho os testes pelo fluxograma e rodo a inferência (etapa 11.3)", "", ""],
     ["12", "Verifico pressupostos e, se preciso, troco pelo teste de postos", "", ""],
     ["13", "Ajusto o modelo e monto a tabela de cinco colunas (etapas 11.4 e 12.4)", "", ""],
     ["14", "Fecho o relatório, o notebook e apresento (AV3)", "", ""],
     ["15", "Prova final", "", ""]],
    [2.0, 9.0, 3.0, 2.0], legenda="Quadro 41 — O meu cronograma")

# ================================================================ ETAPA 15
h1("Etapa 15", "Referências")
entrega("Em todas as entregas.")

caixa_info("O que entra aqui.", [
    "Em ABNT, em ordem alfabética de sobrenome. Entram as obras citadas no texto, e só elas. "
    "As bases de dados também se citam, com o produtor, o nome da tabela ou série, e a ",
    ("data de acesso", True), "."])

caixa_exemplo([
    "GIL, Antonio Carlos. Como elaborar projetos de pesquisa. 7. ed. Barueri: Atlas, 2022.",
    "",
    "INSTITUTO BRASILEIRO DE GEOGRAFIA E ESTATÍSTICA. Demografia das Empresas e "
    "Empreendedorismo: tabela 9949. Rio de Janeiro: IBGE, 2024. Disponível em: "
    "https://sidra.ibge.gov.br/tabela/9949. Acesso em: 14 out. 2026."])

campo("REFERÊNCIAS", linhas=14)

# ================================================================ ETAPA 16
h1("Etapa 16", "Registro de uso de inteligência artificial")
entrega("AV3 — encontro 14. A omissão é falta de honestidade acadêmica.")

caixa_info("A regra do semestre.", [
    ("A IA explica, o código calcula e o estudante confere e escreve.", True),
    (" O assistente é permitido e até incentivado, sobretudo para adaptar a análise à base do "
     "seu projeto. A exigência é de transparência: quem usa, registra. O registro ", False),
    ("não desconta pontos", True),
    ("; a omissão é que é falta.", False)])

h2("16.1", "Os três modos de uso")

tabela(
    ["Modo", "O que você faz", "Como se confere"],
    [["Pergunta", "Pergunta conceitual, com o cartão de contexto colado",
      "A resposta cita um número que está no cartão"],
     ["Encomenda", "Pede uma troca de coluna, um filtro ou um agrupamento",
      "O número bate; os grupos somam o total; o gráfico é o que foi pedido"],
     ["Leitura", "Pede ajuda para ler uma saída do notebook",
      "A frase é sua, e cada número dela está numa saída executada"]],
    [2.2, 6.0, 7.8], legenda="Quadro 42 — Os três modos de uso do assistente")

h2("16.2", "O meu registro")

check([("Não usei assistente de IA neste projeto.", True)])

tabela(
    ["#", "Modo", "O que eu pedi", "A checagem que eu fiz", "O que eu mudei depois"],
    [["1", "", "", "", ""], ["2", "", "", "", ""], ["3", "", "", "", ""],
     ["4", "", "", "", ""], ["5", "", "", "", ""]],
    [0.8, 1.8, 5.0, 4.4, 4.0], legenda="Quadro 43 — Registro de uso de IA")

# ================================================================ ETAPA 17
h1("Etapa 17", "A apresentação de 3 a 5 minutos")
entrega("AV3 — encontro 14. Vale 25 pontos.")

caixa_info("O que entra aqui.", [
    "Um slide — no máximo dois, se o segundo for a Figura 1 — e cinco blocos de cerca de um "
    "minuto. ", ("Quem lê o slide não está apresentando", True),
    ": o slide organiza, a fala comunica. Números grandes, uma figura só e nenhuma tabela de "
    "cinco colunas projetada inteira."])

tabela(
    ["Bloco", "O que dizer", "No exemplo da aula", "O meu"],
    [["1. A pergunta", "O problema de gestão, em uma frase",
      "“O preço varia com a bandeira ou com a região?”", ""],
     ["2. O dado", "Fonte, recorte, unidade e n", "“ANP, 2025, 236 postos do Maranhão”", ""],
     ["3. O achado", "Um número, com a frase completa",
      "“R$ 0,61 a mais no interior; a marca, R$ 0,05”", ""],
     ["4. O limite", "O que o resultado não autoriza",
      "“Não é causa; vale para os postos pesquisados”", ""],
     ["5. A implicação", "O que um gestor faria com isso",
      "“Na compra da frota, a região pesa mais que a bandeira”", ""]],
    [2.4, 3.6, 5.0, 5.0], legenda="Quadro 44 — Os cinco blocos da apresentação")

caixa_info("Como é cronometrado.", [
    "A ordem é sorteada no início do bloco. O professor avisa aos 3 minutos e corta aos 5. "
    "Cada apresentação termina com uma pergunta da turma ou do professor, respondida em até "
    "um minuto. Quem assiste anota, para cada colega, o achado e o limite em uma linha."])

campo("A PERGUNTA QUE EU ESPERO RECEBER, E A MINHA RESPOSTA", linhas=3)

# ================================================================ ANEXO A
h1("Anexo A", "Checklist final, antes de entregar")

caixa_info("Como usar.", [
    "Percorra a lista com o arquivo aberto, e não de memória. ",
    ("A avaliação é sobre o que está no arquivo, e não sobre o que você lembra ter feito.",
     True),
    (" Reabra o link compartilhado e confira.", False)])

h3("O projeto, neste documento")
check([
    "Nenhum campo em branco. O que não se aplica está escrito como “não se aplica”, com o "
    "motivo.",
    "O problema (1.3) está como pergunta e passa pelas seis regras (1.4).",
    "Todo termo ambíguo do problema tem definição operacional (1.5).",
    "Os objetivos usam verbos da coluna ✓ e o geral responde à pergunta (3).",
    "Cada hipótese tem direção, variáveis e um teste que a decide (4 e 11.3).",
    "Toda variável tem nível, unidade e coluna da base (5.2).",
    "A fonte passou pelos seis critérios e pelos quatro defeitos (7.2 e 7.3).",
    "A matriz de amarração não tem nenhum dos quatro buracos (10.3).",
    "As quatro peças de resultado estão presentes, cada uma com a sua frase (12).",
    "As três limitações estão escritas, e a resposta a “o que não autoriza” está pronta (13).",
    "As referências estão em ABNT, e as bases citadas com data de acesso (15).",
    "O registro de IA está preenchido, ou marcado “não usei” (16).",
])

h3("O notebook")
check([
    "A CONFIG aponta para a minha base, e a carga funciona.",
    "Ambiente de execução → Reiniciar e executar tudo roda sem erro, do início ao fim.",
    "A conferência “os grupos somam o total” deu [ok].",
    "Cada número que está neste documento aparece numa saída salva do notebook.",
    "O link foi compartilhado com o professor, e eu reabri o link para conferir.",
])

h3("A apresentação")
check([
    "Entre 3 e 5 minutos, cronometrados em casa pelo menos uma vez.",
    "Um slide, com os cinco blocos (ou dois, se o segundo for a Figura 1).",
    "O número principal vem com a frase completa, não sozinho.",
    "Eu tenho resposta pronta para “o que esse resultado não autoriza dizer?”.",
])

# ================================================================ ANEXO B
h1("Anexo B", "Reprodutibilidade: a CONFIG do notebook")

caixa_info("O que é.", [
    "O notebook do encontro 14 já é um relatório pronto sobre os postos da ANP. Você troca ",
    ("só os seis valores", True),
    (" do bloco de configuração, executa tudo, e cada seção se refaz com a sua base. Anote "
     "aqui os valores que você usou: é isso que torna o seu resultado reproduzível por outra "
     "pessoa.", False)])

tabela(
    ["Parâmetro", "O que é", "No exemplo", "O meu valor"],
    [["ARQUIVO", "A base: “exemplo” ou o nome do seu CSV", "\"exemplo\"", ""],
     ["FONTE", "A citação da sua fonte", "\"ANP, SHPC, 2025\"", ""],
     ["RESPOSTA", "Numérica: o que se quer explicar", "\"preco_gasolina\"", ""],
     ["GRUPO", "Categórica: os grupos do teste", "\"regiao\"", ""],
     ["EXPLICATIVAS", "As que entram no modelo", "[\"com_marca\", \"ilha\"]", ""],
     ["EVENTO", "0/1 para a logística, ou vazio", "\"vende_etanol\"", ""]],
    [2.8, 5.2, 4.0, 4.0], legenda="Quadro 45 — A CONFIG do meu notebook")

caixa_info("A conferência muda de papel.", [
    "Com a base de exemplo, cada número aparece com [ok], porque o gabarito é conhecido. Com "
    "a sua base, aparece [--] e o seu número, para você anotar. A conferência que vale para "
    "qualquer base — ", ("“os grupos somam o total”", True),
    " — continua devolvendo [ok] ou [X]. Se ela der [X], há erro de filtro ou de agrupamento, "
    "e nenhum número posterior vale."])

# ================================================================ ANEXO C
h1("Anexo C", "Como o projeto é avaliado")

para("A rubrica da AV3, encontro 14. Os pontos entre parênteses são os de cada item.",
     after=6)

tabela(
    ["Peça", "O que se avalia", "Pontos"],
    [["Relatório",
      "Pergunta, método e níveis coerentes (10) · Tabela 1 e Figura 1 com n, unidade e fonte "
      "(10) · teste escolhido pelo fluxograma e bem lido (10) · modelo em cinco colunas, sem "
      "verbo causal (10) · limitações (5) · frases completas (5)", "50"],
     ["Notebook",
      "Executa do início ao fim sem erro (10) · CONFIG e base do projeto carregando (5) · "
      "conferências executadas (5) · registro de IA preenchido ou “não usei” (5)", "25"],
     ["Apresentação",
      "Entre 3 e 5 minutos (5) · os cinco blocos (10) · o número principal com a frase "
      "completa (5) · resposta a uma pergunta (5)", "25"]],
    [2.6, 11.4, 2.0], legenda="Quadro 46 — Rubrica da AV3")

caixa_info("O que separa a nota cheia da parcial.", [
    "O número certo com a frase que passa no teste das quatro perguntas. ",
    ("Número sem frase vale menos da metade", True),
    (", e verbo causal sobre dado observacional custa os pontos do modelo.", False)])

# ================================================================ ANEXO D
h1("Anexo D", "Glossário mínimo")

tabela(
    ["Termo", "O que é, em uma frase"],
    [["População", "Todos os casos sobre os quais se quer concluir"],
     ["Amostra", "O subconjunto da população que foi efetivamente observado"],
     ["Unidade de observação", "O que é um caso: uma empresa, um posto, um município, um mês"],
     ["Parâmetro", "O número da população, em geral desconhecido (μ, σ, π)"],
     ["Estimativa", "O número calculado na amostra, que estima o parâmetro (x̄, s, p)"],
     ["Variável", "Característica que varia entre os casos"],
     ["Nível de mensuração", "Nominal, ordinal, intervalar ou de razão; decide o que se "
      "pode calcular"],
     ["Definição operacional", "A que diz como o fenômeno é medido"],
     ["Hipótese", "Afirmação sobre a relação entre duas variáveis, feita antes de olhar o "
      "resultado"],
     ["Hipótese nula (H₀)", "A que afirma ausência de diferença ou de associação"],
     ["Nível de significância (α)", "O risco de rejeitar a nula sendo ela verdadeira, "
      "declarado antes"],
     ["Valor-p", "A probabilidade de ver um resultado tão extremo quanto o observado, se a "
      "nula fosse verdadeira"],
     ["Erro tipo I", "Rejeitar a nula sendo ela verdadeira"],
     ["Erro tipo II", "Não rejeitar a nula sendo ela falsa"],
     ["Margem de erro", "O quanto a estimativa pode se afastar do parâmetro, no nível de "
      "confiança declarado"],
     ["Intervalo de confiança", "A faixa de valores compatíveis com os dados, no nível "
      "declarado"],
     ["Erro padrão", "O desvio padrão da distribuição amostral da estimativa"],
     ["Coeficiente de variação", "Desvio padrão dividido pela média; permite comparar "
      "dispersões de escalas diferentes"],
     ["Correlação", "Medida de associação linear entre duas variáveis numéricas; não é causa"],
     ["R²", "Fração da variação da resposta que o modelo acompanha; não é prova de "
      "mecanismo"],
     ["Razão de chances (odds ratio)", "Quanto a chance do evento se multiplica por unidade "
      "da explicativa, na logística"],
     ["Variável de confusão", "A que afeta resposta e explicativa ao mesmo tempo e explica "
      "parte da associação observada"],
     ["Matriz de amarração", "Tabela que liga problema, objetivo, hipótese, variáveis, base e "
      "técnica"]],
    [4.5, 11.5], legenda="Quadro 47 — Glossário mínimo do semestre")

espaco(10)
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "Métodos e Técnicas de Pesquisa Quantitativa  ·  Administração (Bacharelado)  ·  UFMA",
    size=9, bold=True, color=VINHO, font=F_TIT)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "Prof. Dr. Tadeu Gomes Teixeira  ·  material didático de uso livre na disciplina",
    size=9, color=CINZA, font=F_TIT)

import sys

out = os.path.join(RAIZ, "Template_Projeto_Pesquisa_Quantitativa.docx")
if len(sys.argv) > 1:
    out = sys.argv[1]
doc.save(out)
print("salvo:", out)
